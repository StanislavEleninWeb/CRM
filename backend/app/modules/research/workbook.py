"""Reading uploaded spreadsheets safely.

Formulas are never evaluated: only values cached in the file are read. Macros and
external links are refused or ignored, and size limits apply before parsing.
"""

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

import openpyxl

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200
MAX_ROWS = 20_000
MAX_COLUMNS = 100
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class WorkbookError(ValueError):
    """The file cannot be imported. The message is safe to show to the user."""


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[dict[str, Any]]
    formula_cells_without_cache: int = 0


@dataclass
class ParsedWorkbook:
    kind: str
    leads: Sheet
    shortlist: Sheet | None = None
    summary: list[list[str]] = field(default_factory=list)
    summary_sheet_name: str | None = None
    ignored_sheets: list[str] = field(default_factory=list)


def _normal(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", name.lower()).strip()


def _check_archive(data: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise WorkbookError("This is not a valid .xlsx file.") from exc
    total = 0
    for info in archive.infolist():
        name = info.filename.lower()
        if "vbaproject" in name or (name.endswith(".bin") and "macro" in name):
            raise WorkbookError("Workbooks with macros are not accepted. Save it as .xlsx without macros.")
        total += info.file_size
        if (
            info.compress_size
            and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO
            and info.file_size > 1_000_000
        ):
            raise WorkbookError("The workbook is compressed in a way that is not accepted.")
    if total > MAX_UNCOMPRESSED_BYTES:
        raise WorkbookError("The workbook is too large once unpacked.")


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.replace("\x00", "")
    return value


def _read_sheet(values_ws: Any, formulas_ws: Any) -> Sheet:
    rows_iter = values_ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return Sheet(name=values_ws.title, headers=[], rows=[])
    if len(header_row) > MAX_COLUMNS:
        raise WorkbookError(f"A sheet has more than {MAX_COLUMNS} columns.")
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    formula_iter = formulas_ws.iter_rows(min_row=2, values_only=True)
    rows: list[dict[str, Any]] = []
    missing_cache = 0
    for index, values in enumerate(rows_iter, start=2):
        if index > MAX_ROWS + 1:
            raise WorkbookError(f"A sheet has more than {MAX_ROWS} rows.")
        formulas = next(formula_iter, ())
        for position, value in enumerate(values):
            source = formulas[position] if position < len(formulas) else None
            if value is None and isinstance(source, str) and source.startswith("="):
                missing_cache += 1
        if all(v is None or v == "" for v in values):
            continue
        row = {headers[i]: _cell(values[i]) for i in range(min(len(headers), len(values))) if headers[i]}
        row["__row__"] = index
        rows.append(row)
    return Sheet(
        name=values_ws.title, headers=[h for h in headers if h], rows=rows, formula_cells_without_cache=missing_cache
    )


def read_xlsx(data: bytes) -> ParsedWorkbook:
    if len(data) > MAX_FILE_BYTES:
        raise WorkbookError(f"Files are limited to {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    _check_archive(data)
    try:
        values = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        formulas = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError("The workbook could not be read.") from exc

    leads_ws = shortlist_ws = summary_ws = None
    ignored: list[str] = []
    for ws in values.worksheets:
        name = _normal(ws.title)
        if leads_ws is None and (("all" in name and "lead" in name) or name in ("leads", "prospects")):
            leads_ws = ws
        elif shortlist_ws is None and "shortlist" in name:
            shortlist_ws = ws
        elif summary_ws is None and "summary" in name:
            summary_ws = ws
        else:
            ignored.append(ws.title)
    if leads_ws is None:
        # Fall back to the first sheet that has a lead identifier or a business name column.
        for ws in values.worksheets:
            first = next(ws.iter_rows(max_row=1, values_only=True), ())
            names = {_normal(str(h)) for h in first if h}
            if names & {"lead id", "business name", "company", "company name", "name"}:
                leads_ws = ws
                ignored = [n for n in ignored if n != ws.title]
                break
    if leads_ws is None:
        raise WorkbookError("No sheet with a lead list was found.")

    result = ParsedWorkbook(kind="xlsx", leads=_read_sheet(leads_ws, formulas[leads_ws.title]), ignored_sheets=ignored)
    if shortlist_ws is not None:
        result.shortlist = _read_sheet(shortlist_ws, formulas[shortlist_ws.title])
    if summary_ws is not None:
        result.summary_sheet_name = summary_ws.title
        for row in summary_ws.iter_rows(max_row=300, max_col=6, values_only=True):
            if any(cell is not None for cell in row):
                result.summary.append(["" if cell is None else str(cell)[:1000] for cell in row])
    return result


def read_csv(data: bytes) -> ParsedWorkbook:
    if len(data) > MAX_FILE_BYTES:
        raise WorkbookError(f"Files are limited to {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise WorkbookError("CSV files must be UTF-8 encoded.") from exc
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        header_row = next(reader)
    except StopIteration as exc:
        raise WorkbookError("The file is empty.") from exc
    if len(header_row) > MAX_COLUMNS:
        raise WorkbookError(f"The file has more than {MAX_COLUMNS} columns.")
    headers = [h.strip() for h in header_row]
    rows: list[dict[str, Any]] = []
    for index, values in enumerate(reader, start=2):
        if index > MAX_ROWS + 1:
            raise WorkbookError(f"The file has more than {MAX_ROWS} rows.")
        if not any(v.strip() for v in values):
            continue
        row: dict[str, Any] = {
            headers[i]: values[i].replace("\x00", "") for i in range(min(len(headers), len(values))) if headers[i]
        }
        row["__row__"] = index
        rows.append(row)
    return ParsedWorkbook(kind="csv", leads=Sheet(name="CSV", headers=[h for h in headers if h], rows=rows))


def read_upload(filename: str, data: bytes) -> ParsedWorkbook:
    lower = filename.lower()
    if lower.endswith(".xlsm") or lower.endswith(".xls"):
        raise WorkbookError("Only .xlsx and .csv files are accepted.")
    if lower.endswith(".xlsx"):
        return read_xlsx(data)
    if lower.endswith(".csv"):
        return read_csv(data)
    raise WorkbookError("Only .xlsx and .csv files are accepted.")
