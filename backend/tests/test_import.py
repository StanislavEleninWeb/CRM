"""Prospect import and export: synthetic cases always run; reference-workbook cases need the private fixture."""

import io
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.modules.research.exporter import LEAD_HEADERS, SHORTLIST_HEADERS
from app.modules.research.tasks import commit_import_task
from app.modules.research.workbook import MAX_FILE_BYTES
from tests.helpers import API, create_workspace, join, sign_in

HEADERS = [h.replace("{tenant}", "SEWEB") for h in LEAD_HEADERS]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def ok(response: Any, status: int = 200) -> Any:
    assert response.status_code == status, response.text
    return response.json() if response.content else None


@pytest.fixture
def owner(make_client: Any) -> TestClient:
    client = make_client()
    sign_in(client, "owner@example.test")
    create_workspace(client, "SEWEB")
    return client


@pytest.fixture
def second_tenant(make_client: Any) -> TestClient:
    client = make_client()
    sign_in(client, "other@example.test")
    create_workspace(client, "SEWEB")
    return client


def lead_row(lead_id: str | None, name: str, **overrides: Any) -> dict[str, Any]:
    row: dict[str, Any] = dict.fromkeys(HEADERS)
    row.update(
        {
            "Lead ID": lead_id,
            "Business name": name,
            "Industry / business type": "Hair & beauty salon",
            "City": "Sofia",
            "Country": "Bulgaria",
            "Website URL": "https://example-salon.bg/",
            "Google Business Profile / Maps URL": "https://www.google.com/maps/place/?q=place_id:ChIJabcdefghijklmnop",
            "Public business phone": "+359 2 111 2233; 0888 111 222",
            "Public business email": "office@example-salon.bg",
            "Contact page URL": "Not found",
            "Other public business contact channel": "facebook.com/examplesalon",
            "Website status": "Loads (HTTPS)",
            "Specific observed issue or opportunity": "Homepage shows template filler text.",
            "Evidence / relevant page URL": "https://example-salon.bg/",
            "Opportunity hypothesis requiring confirmation": "Booking is done by phone only.",
            "Recommended SEWEB service": "Website refresh; online booking",
            "Why this business is a suitable cold prospect": "Active salon with many reviews.",
            "Suggested business benefit": "Clients can book at any hour.",
            "Personalized outreach opening": "Здравейте, разгледах сайта Ви.",
            "Suggested discovery question": "How are appointments scheduled today?",
            "Recommended contact channel": "Email, then phone",
            "Priority score, 0–100": 89,
            "Priority tier: A, B, or C": "A",
            "Evidence confidence: high, medium, or low": "high",
            "Date checked": __import__("datetime").datetime(2026, 10, 7),
            "Outreach status": "Not contacted",
            "Notes": "Checked twice.",
            "Score: evidence strength (0–30)": 28,
            "Score: relevance to SEWEB (0–25)": 23,
            "Score: business value (0–20)": 15,
            "Score: reachability (0–15)": 14,
            "Score: business activity (0–10)": 9,
            "Industry group": "Beauty & wellness",
            "Service category": "Online booking system + AI assistant",
            "Observed opportunity tags": "No online booking / phone- or form-only intake; Other",
        }
    )
    for key, value in overrides.items():
        assert key in row, key
        row[key] = value
    return row


def workbook_bytes(rows: list[dict[str, Any]], shortlist: list[str] | None = None, summary: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All qualified leads"
    ws.append(HEADERS)
    for row in rows:
        ws.append([row[h] for h in HEADERS])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # text stays text even when it starts with "="
    if shortlist is not None:
        sl = wb.create_sheet("Tomorrow’s outreach shortlist")
        sl.append([h.replace("{tenant}", "SEWEB") for h in SHORTLIST_HEADERS])
        for rank, lead_id in enumerate(shortlist, start=1):
            sl.append([rank, lead_id])
    if summary:
        sm = wb.create_sheet("Research summary")
        sm.append(["Synthetic research summary"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def upload(client: TestClient, data: bytes, name: str = "prospects.xlsx", content_type: str = XLSX) -> Any:
    return client.post(f"{API}/imports", files={"file": (name, io.BytesIO(data), content_type)})


def import_and_commit(client: TestClient, data: bytes, name: str = "prospects.xlsx") -> dict[str, Any]:
    created = ok(upload(client, data, name), 202)
    assert created["status"] == "ready", created
    committed = ok(client.post(f"{API}/imports/{created['id']}/commit"), 202)
    assert committed["status"] == "committed", committed
    return committed  # type: ignore[no-any-return]


def tenant_query(engine: Any, client: TestClient, sql: str, **params: Any) -> list[Any]:
    tenant_id = ok(client.get(f"{API}/tenant"))["id"]
    with engine.begin() as conn:
        conn.execute(text("SELECT set_config('app.tenant_id', :t, true)"), {"t": tenant_id})
        return list(conn.execute(text(sql), {"t": tenant_id, **params}).all())


# --- synthetic -------------------------------------------------------------------------------


def test_preview_changes_nothing_until_commit(owner: TestClient) -> None:
    data = workbook_bytes(
        [lead_row("T-001", "Салон Аврора"), lead_row("T-002", "Vet Clinic", City="Varna")], shortlist=["T-002", "T-001"]
    )
    created = ok(upload(owner, data), 202)
    report = created["report"]
    assert created["status"] == "ready"
    assert report["counts"] == {"create": 2, "rows": 2, "with_errors": 0, "with_warnings": 0}
    assert len(report["mapping"]) == 35 and report["unmapped_headers"] == []
    assert report["tiers"] == {"A": 2} and report["confidence"] == {"high": 2}
    assert report["missing_not_found"] == {"contact_page": 2}
    assert report["shortlist"] == {
        "sheet": "Tomorrow’s outreach shortlist",
        "rows": 2,
        "unresolved_lead_ids": [],
        "creates_leads": False,
        "formula_cells_without_cache": 0,
    }
    assert ok(owner.get(f"{API}/leads"))["total"] == 0 and ok(owner.get(f"{API}/companies"))["total"] == 0
    rows = ok(owner.get(f"{API}/imports/{created['id']}/rows"))
    assert [
        (r["external_id"], r["business_name"], r["action"], r["score_total"], r["tier"]) for r in rows["items"]
    ] == [("T-001", "Салон Аврора", "create", 89, "A"), ("T-002", "Vet Clinic", "create", 89, "A")]

    committed = ok(owner.post(f"{API}/imports/{created['id']}/commit"), 202)
    assert committed["status"] == "committed" and committed["result"]["created"] == 2
    assert ok(owner.get(f"{API}/leads"))["total"] == 2
    assert owner.post(f"{API}/imports/{created['id']}/commit").status_code == 409


def test_imported_row_becomes_structured_records(owner: TestClient, migrator_engine: Any) -> None:
    import_and_commit(owner, workbook_bytes([lead_row("T-001", "Салон Аврора")]))
    lead = ok(owner.get(f"{API}/leads"))["items"][0]
    assert (
        lead["external_id"] == "T-001" and lead["status"] == "qualified" and lead["outreach_status"] == "not_contacted"
    )
    assert lead["source"] == "import" and lead["company_name"] == "Салон Аврора"
    company = ok(owner.get(f"{API}/companies/{lead['company_id']}"))
    assert company["domain"] == "example-salon.bg" and company["industry_group"] == "Beauty & wellness"
    channels = {(c["kind"], c["raw_value"]): c for c in company["channels"]}
    assert set(channels) == {
        ("phone", "+359 2 111 2233"),
        ("phone", "0888 111 222"),
        ("email", "office@example-salon.bg"),
        ("social", "facebook.com/examplesalon"),
    }
    assert channels[("phone", "0888 111 222")]["normalized_value"] == "+359888111222"
    assert all(c["source_type"] == "user_import" and c["source_date"] == "2026-10-07" for c in channels.values())
    assert all(c["verification_state"] == "unverified" for c in channels.values())  # imported is not verified

    a = tenant_query(migrator_engine, owner, "SELECT * FROM lead_assessments WHERE tenant_id = :t")[0]
    assert str(a.checked_on) == "2026-10-07" and a.confidence == "high"
    assert (a.source_type, a.verification_state) == ("user_import", "unverified")
    assert (a.website_base, a.website_transport, a.website_status_raw) == ("loads", "https", "Loads (HTTPS)")
    assert (a.listing_id_type, a.listing_id) == ("place_id", "ChIJabcdefghijklmnop")
    assert a.contact_states == {"contact_page": {"state": "not_found", "raw": "Not found"}}
    assert (a.preferred_channel, a.fallback_channels, a.channel_recommendation_raw) == (
        "email",
        ["phone"],
        "Email, then phone",
    )
    assert a.recommended_services == ["Website refresh", "online booking"]
    assert a.outreach_opening == "Здравейте, разгледах сайта Ви."
    assert len(a.raw_row) == 35 and a.raw_row["Public business phone"] == "+359 2 111 2233; 0888 111 222"
    observations = tenant_query(
        migrator_engine, owner, "SELECT text, evidence_url, verification_state FROM observations WHERE tenant_id = :t"
    )
    assert observations == [("Homepage shows template filler text.", "https://example-salon.bg/", "unverified")]
    hypotheses = tenant_query(migrator_engine, owner, "SELECT text, status FROM hypotheses WHERE tenant_id = :t")
    assert hypotheses == [("Booking is done by phone only.", "unconfirmed")]  # kept apart from observations
    score = ok(owner.get(f"{API}/leads/{lead['id']}/scores"))[0]
    assert (score["total"], score["tier"], score["origin"], score["rubric_version"], score["source_total"]) == (
        89,
        "A",
        "import",
        1,
        89,
    )
    tags = tenant_query(
        migrator_engine,
        owner,
        "SELECT t.name::text, tg.origin FROM taggings tg JOIN tags t ON t.id = tg.tag_id WHERE tg.tenant_id = :t ORDER BY 1",
    )
    assert tags == [("No online booking / phone- or form-only intake", "import"), ("Other", "import")]
    assert tenant_query(migrator_engine, owner, "SELECT name FROM service_catalog WHERE tenant_id = :t") == [
        ("Online booking system + AI assistant",)
    ]


def test_reimport_updates_research_without_duplicating_or_undoing_user_work(
    owner: TestClient, migrator_engine: Any
) -> None:
    rows = [lead_row("T-001", "Salon Aurora"), lead_row("T-002", "Vet Clinic")]
    import_and_commit(owner, workbook_bytes(rows, shortlist=["T-001", "T-002"]))
    lead = next(item for item in ok(owner.get(f"{API}/leads"))["items"] if item["external_id"] == "T-001")
    company = ok(owner.get(f"{API}/companies/{lead['company_id']}"))
    phone = next(c for c in company["channels"] if c["raw_value"] == "0888 111 222")
    ok(owner.patch(f"{API}/channels/{phone['id']}", json={"do_not_contact": True, "restriction_reason": "Opted out"}))
    ok(owner.patch(f"{API}/leads/{lead['id']}", json={"outreach_status": "called"}))
    ok(owner.patch(f"{API}/companies/{lead['company_id']}", json={"city": "Plovdiv"}))
    ok(
        owner.post(
            f"{API}/leads/{lead['id']}/scores",
            json={
                "components": {"evidence": 10, "relevance": 10, "value": 10, "reachability": 10, "activity": 5},
                "reason": "Site was rebuilt last week",
            },
        ),
        201,
    )

    # Identical file again: nothing new.
    again = ok(upload(owner, workbook_bytes(rows, shortlist=["T-001", "T-002"])), 202)
    assert again["report"]["counts"]["update"] == 2 and "create" not in again["report"]["counts"]
    result = ok(owner.post(f"{API}/imports/{again['id']}/commit"), 202)["result"]
    assert result["unchanged"] == 2 and "created" not in result
    assert ok(owner.get(f"{API}/leads"))["total"] == 2 and ok(owner.get(f"{API}/companies"))["total"] == 2
    assert tenant_query(migrator_engine, owner, "SELECT count(*) FROM lead_assessments WHERE tenant_id = :t")[0][0] == 2
    assert tenant_query(migrator_engine, owner, "SELECT count(*) FROM shortlists WHERE tenant_id = :t")[0][0] == 1
    assert tenant_query(migrator_engine, owner, "SELECT count(*) FROM contact_channels WHERE tenant_id = :t")[0][0] == 8

    # Changed research for the same lead: a new snapshot, the old one kept.
    rows[0]["Specific observed issue or opportunity"] = "Homepage was rebuilt; booking link is now broken."
    rows[0]["Score: evidence strength (0–30)"] = 20
    rows[0]["Priority score, 0–100"] = 81
    rows[0]["Public business phone"] = "+359 2 111 2233; 0888 111 222; 0899 555 666"
    result = import_and_commit(owner, workbook_bytes(rows))["result"]
    assert (result["updated"], result["unchanged"], result["override_kept"]) == (1, 1, 1)
    assessments = tenant_query(
        migrator_engine,
        owner,
        "SELECT is_current FROM lead_assessments WHERE tenant_id = :t AND lead_id = :l ORDER BY created_at",
        l=lead["id"],
    )
    assert assessments == [(False,), (True,)]
    observations = tenant_query(
        migrator_engine,
        owner,
        "SELECT text, superseded_at IS NULL FROM observations WHERE tenant_id = :t AND lead_id = :l ORDER BY created_at",
        l=lead["id"],
    )
    assert observations == [
        ("Homepage shows template filler text.", False),
        ("Homepage was rebuilt; booking link is now broken.", True),
    ]

    # The user's work survives: restriction, outreach status, edited city, and the reviewed score.
    company = ok(owner.get(f"{API}/companies/{lead['company_id']}"))
    assert company["city"] == "Plovdiv"
    blocked = next(c for c in company["channels"] if c["raw_value"] == "0888 111 222")
    assert blocked["do_not_contact"] and blocked["dial_uri"] is None
    assert any(c["raw_value"] == "0899 555 666" for c in company["channels"])
    assert ok(owner.get(f"{API}/leads/{lead['id']}"))["outreach_status"] == "called"
    scores = ok(owner.get(f"{API}/leads/{lead['id']}/scores"))
    assert [(s["origin"], s["total"], s["is_current"]) for s in scores] == [
        ("import", 81, False),
        ("override", 45, True),
        ("import", 89, False),
    ]
    timeline = [a["kind"] for a in ok(owner.get(f"{API}/activities", params={"lead_id": lead["id"]}))["items"]]
    assert timeline.count("lead.imported") == 1 and "lead.research_updated" in timeline


def test_conflicts_and_invalid_rows_are_reported_and_skipped(owner: TestClient) -> None:
    rows = [
        lead_row("T-001", "Good Co"),
        lead_row("T-001", "Same ID Again"),
        lead_row("T-003", "Over Cap", **{"Score: evidence strength (0–30)": 45}),
        lead_row("T-004", "Wrong Total", **{"Priority score, 0–100": 95, "Priority tier: A, B, or C": "B"}),
        lead_row("T-005", None),
        lead_row("T-006", "Partial Score", **{"Score: business activity (0–10)": None}),
        lead_row(
            "T-007", "Unscored", **{h: None for h in HEADERS if h.startswith("Score:") or h.startswith("Priority")}
        ),
        lead_row(
            "T-008",
            "Bad Details",
            **{
                "Public business email": "not-an-email",
                "Website URL": "see facebook",
                "Evidence confidence: high, medium, or low": "certain",
                "Date checked": "last week",
            },
        ),
    ]
    created = ok(upload(owner, workbook_bytes(rows, shortlist=["T-001", "T-404"])), 202)
    report = created["report"]
    assert report["counts"]["rows"] == 8 and report["counts"]["with_errors"] == 4 and report["counts"]["skip"] == 4
    assert report["score_mismatches"] == 1 and report["tiers"]["unscored"] == 3
    assert report["shortlist"]["unresolved_lead_ids"] == ["T-404"]
    by_row = {r["business_name"]: r for r in ok(owner.get(f"{API}/imports/{created['id']}/rows"))["items"]}
    assert "appears more than once" in by_row["Same ID Again"]["issues"][0]["message"]
    assert "between 0 and 30" in by_row["Over Cap"]["issues"][0]["message"]
    assert "partial assessment" in by_row["Partial Score"]["issues"][0]["message"]
    assert {i["field"] for i in by_row["Wrong Total"]["issues"]} == {"score_total", "tier"}
    assert by_row["Wrong Total"]["score_total"] == 89 and by_row["Wrong Total"]["action"] == "create"
    assert {i["field"] for i in by_row["Bad Details"]["issues"]} == {
        "email",
        "website_url",
        "confidence",
        "date_checked",
    }
    issues_only = ok(owner.get(f"{API}/imports/{created['id']}/rows", params={"only": "issues"}))
    assert issues_only["total"] == 6
    bad = next(r for r in issues_only["items"] if r["business_name"] == "Over Cap")
    assert owner.patch(f"{API}/imports/{created['id']}/rows/{bad['id']}", json={"action": "create"}).status_code == 422

    result = ok(owner.post(f"{API}/imports/{created['id']}/commit"), 202)["result"]
    assert (result["created"], result["skipped"]) == (4, 4)
    leads = {item["external_id"]: item for item in ok(owner.get(f"{API}/leads"))["items"]}
    assert set(leads) == {"T-001", "T-004", "T-007", "T-008"}
    assert ok(owner.get(f"{API}/leads/{leads['T-007']['id']}/scores")) == []  # unscored, not zero
    assert (
        ok(owner.get(f"{API}/leads/{leads['T-004']['id']}/scores"))[0]["total"] == 89
    )  # computed, not the source's 95
    bad_company = ok(owner.get(f"{API}/companies/{leads['T-008']['company_id']}"))
    assert bad_company["website_url"] is None and not any(c["kind"] == "email" for c in bad_company["channels"])


def test_formula_injection_text_is_kept_as_text_in_and_out(owner: TestClient) -> None:
    hostile = {
        "Business name": '=HYPERLINK("http://evil.example","Click")',
        "Notes": "+cmd|' /C calc'!A0",
        "Suggested discovery question": "@SUM(1+1)",
        "Personalized outreach opening": "-2+3 е добър резултат",
        "Public business phone": "+359 2 900 0001",
    }
    import_and_commit(owner, workbook_bytes([lead_row("T-001", "x", **hostile)]))
    company = ok(owner.get(f"{API}/companies"))["items"][0]
    assert company["name"] == hostile["Business name"]

    exported = owner.get(f"{API}/exports/prospects.xlsx")
    assert exported.status_code == 200 and exported.headers["content-type"] == XLSX
    assert exported.headers["x-content-type-options"] == "nosniff"
    sheet = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=False)["All qualified leads"]
    header = {cell.value: cell.column for cell in sheet[1]}
    for column, value in hostile.items():
        cell = sheet.cell(row=2, column=header[column])
        assert cell.value == value, column  # unchanged, including the ordinary phone number
        assert cell.data_type == "s", column  # a string cell, never a formula
    raw_xml = zipfile.ZipFile(io.BytesIO(exported.content)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in raw_xml and "<f " not in raw_xml
    assert "export.prospects" in [e["action"] for e in ok(owner.get(f"{API}/audit-events"))["items"]]


def test_export_shows_the_current_status_and_channel_not_the_imported_text(owner: TestClient) -> None:
    import_and_commit(owner, workbook_bytes([lead_row("T-001", "Salon Aurora"), lead_row("T-002", "Vet Clinic")]))
    leads = {item["external_id"]: item for item in ok(owner.get(f"{API}/leads"))["items"]}

    def exported() -> dict[str, dict[str, Any]]:
        sheet = openpyxl.load_workbook(io.BytesIO(owner.get(f"{API}/exports/prospects.xlsx").content))[
            "All qualified leads"
        ]
        header = [c.value for c in sheet[1]]
        return {r[0]: dict(zip(header, r, strict=True)) for r in sheet.iter_rows(min_row=2, values_only=True)}

    before = exported()
    assert before["T-001"]["Outreach status"] == "Not contacted"  # unchanged: the imported wording
    assert before["T-001"]["Recommended contact channel"] == "Email, then phone"

    detail = ok(owner.get(f"{API}/prospects/{leads['T-001']['id']}"))
    phone = next(c for c in detail["channels"] if c["kind"] == "phone")
    call = ok(owner.post(f"{API}/prospects/{leads['T-001']['id']}/calls", json={"channel_id": phone["id"]}), 201)
    ok(owner.post(f"{API}/calls/{call['id']}/outcome", json={"outcome": "connected"}))
    ok(
        owner.patch(
            f"{API}/prospects/{leads['T-001']['id']}/assessment",
            json={"preferred_channel": "phone", "channel_instruction": "ask for the owner"},
        )
    )
    after = exported()
    assert after["T-001"]["Outreach status"] == "Contacted"
    assert after["T-001"]["Recommended contact channel"] == "Phone (ask for the owner)"
    assert after["T-002"]["Outreach status"] == "Not contacted"
    assert after["T-002"]["Recommended contact channel"] == "Email, then phone"


def test_formulas_are_never_evaluated_and_missing_caches_are_reported(owner: TestClient) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "All qualified leads"
    ws.append(HEADERS)
    row = lead_row("T-001", "Formula Co")
    ws.append([row[h] for h in HEADERS])
    total_col, tier_col = HEADERS.index("Priority score, 0–100") + 1, HEADERS.index("Priority tier: A, B, or C") + 1
    ws.cell(row=2, column=total_col).value = "=SUM(AB2:AF2)"  # saved by a tool that stores no cached result
    ws.cell(row=2, column=tier_col).value = '=IF(V2>=80,"A","B")'
    ws.cell(row=2, column=HEADERS.index("Notes") + 1).value = "=1+1"
    buffer = io.BytesIO()
    wb.save(buffer)
    created = ok(upload(owner, buffer.getvalue()), 202)
    assert created["status"] == "ready" and created["report"]["formula_cells_without_cache"] == 3
    assert created["report"]["tiers"] == {"A": 1}  # computed from the components, not from the formula
    import_and_commit(owner, buffer.getvalue())
    lead = ok(owner.get(f"{API}/leads"))["items"][0]
    assert ok(owner.get(f"{API}/leads/{lead['id']}/scores"))[0]["total"] == 89


def test_unsafe_or_unreadable_files_are_refused(owner: TestClient) -> None:
    assert upload(owner, b"x", "list.xlsm").status_code == 422
    assert upload(owner, b"x", "list.xls").status_code == 422
    assert upload(owner, b"x", "list.pdf").status_code == 422
    assert upload(owner, b"", "empty.xlsx").status_code == 422
    assert upload(owner, b"x" * (MAX_FILE_BYTES + 1), "huge.xlsx").status_code == 422

    garbage = ok(upload(owner, b"this is not a zip archive", "broken.xlsx"), 202)
    assert garbage["status"] == "failed" and garbage["error"] == "This is not a valid .xlsx file."

    macro = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(workbook_bytes([lead_row("T-1", "A")]))) as source,
        zipfile.ZipFile(macro, "w") as target,
    ):
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"macro")
    with_macro = ok(upload(owner, macro.getvalue(), "renamed.xlsx"), 202)
    assert with_macro["status"] == "failed" and "macros" in with_macro["error"]

    no_leads = openpyxl.Workbook()
    no_leads.active.title = "Notes"
    no_leads.active.append(["Colour", "Shape"])
    buffer = io.BytesIO()
    no_leads.save(buffer)
    assert ok(upload(owner, buffer.getvalue(), "other.xlsx"), 202)["error"] == "No sheet with a lead list was found."
    assert owner.post(f"{API}/imports/{garbage['id']}/commit").status_code == 409
    assert ok(owner.get(f"{API}/leads"))["total"] == 0


def test_csv_import_with_review_of_possible_duplicates(owner: TestClient) -> None:
    import_and_commit(owner, workbook_bytes([lead_row("T-001", "Салон Аврора")]))
    csv_text = (
        "Business name,City,Country,Public business phone,Score: evidence strength (0–30),Extra column\r\n"
        "Салон Аврора,Sofia,Bulgaria,0888 999 000,,keep me\r\n"
        "Нова Фирма,Varna,Bulgaria,Not found,,\r\n"
    )
    created = ok(upload(owner, csv_text.encode("utf-8-sig"), "more.csv", "text/csv"), 202)
    assert created["status"] == "ready" and created["report"]["unmapped_headers"] == ["Extra column"]
    assert created["report"]["counts"]["skip"] == 1 and created["report"]["counts"]["create"] == 1
    duplicates = ok(owner.get(f"{API}/imports/{created['id']}/rows", params={"only": "duplicates"}))["items"]
    assert (
        len(duplicates) == 1
        and duplicates[0]["duplicate_reason"] == "same_name_and_city"
        and duplicates[0]["action"] == "skip"
    )
    # The reviewer decides it is a different business after all.
    ok(owner.patch(f"{API}/imports/{created['id']}/rows/{duplicates[0]['id']}", json={"action": "create"}))
    assert ok(owner.post(f"{API}/imports/{created['id']}/commit"), 202)["result"]["created"] == 2
    assert ok(owner.get(f"{API}/leads"))["total"] == 3
    assert ok(owner.get(f"{API}/companies", params={"q": "нова"}))["total"] == 1  # Cyrillic survives and is searchable
    assert (
        upload(owner, "Business name\n\xff\xfe".encode("latin-1"), "bad.csv", "text/csv").json()["error"]
        == "CSV files must be UTF-8 encoded."
    )


def test_column_mapping_can_be_corrected_before_commit(owner: TestClient) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Company", "Town", "Telephone number"])
    ws.append(["Acme Bakery", "Ruse", "+359 82 123 456"])
    buffer = io.BytesIO()
    wb.save(buffer)
    created = ok(upload(owner, buffer.getvalue()), 202)
    assert created["report"]["mapping"] == {"business_name": "Company", "phone": "Telephone number"}
    assert created["report"]["unmapped_headers"] == ["Town"]
    assert (
        owner.put(f"{API}/imports/{created['id']}/mapping", json={"mapping": {"city": "No such column"}}).status_code
        == 422
    )
    assert owner.put(f"{API}/imports/{created['id']}/mapping", json={"mapping": {"planet": "Town"}}).status_code == 422
    remapped = ok(owner.put(f"{API}/imports/{created['id']}/mapping", json={"mapping": {"city": "Town"}}), 202)
    assert remapped["status"] == "ready" and remapped["report"]["mapping"]["city"] == "Town"
    ok(owner.post(f"{API}/imports/{created['id']}/commit"), 202)
    assert ok(owner.get(f"{API}/companies"))["items"][0]["city"] == "Ruse"


def test_import_and_export_need_permission_and_stay_inside_the_tenant(
    owner: TestClient, second_tenant: TestClient, make_client: Any
) -> None:
    rep = make_client()
    sign_in(rep, "rep@example.test")
    join(owner, rep, "rep@example.test", "representative")
    data = workbook_bytes([lead_row("T-001", "Salon Aurora")])
    assert upload(rep, data).status_code == 403
    assert rep.get(f"{API}/exports/prospects.xlsx").status_code == 403
    created = import_and_commit(owner, data)
    for path in (f"/imports/{created['id']}", f"/imports/{created['id']}/rows"):
        assert second_tenant.get(f"{API}{path}").status_code == 404
    assert second_tenant.post(f"{API}/imports/{created['id']}/commit").status_code == 404
    assert ok(second_tenant.get(f"{API}/imports"))["total"] == 0
    theirs = openpyxl.load_workbook(io.BytesIO(second_tenant.get(f"{API}/exports/prospects.xlsx").content))
    assert theirs["All qualified leads"].max_row == 1  # headers only
    # The same Lead ID in another tenant is a different lead.
    import_and_commit(second_tenant, data)
    assert ok(second_tenant.get(f"{API}/leads"))["total"] == 1 and ok(owner.get(f"{API}/leads"))["total"] == 1


def test_duplicate_task_delivery_does_not_apply_an_import_twice(owner: TestClient, migrator_engine: Any) -> None:
    created = import_and_commit(owner, workbook_bytes([lead_row("T-001", "Salon Aurora")]))
    tenant_id, user_id = ok(owner.get(f"{API}/tenant"))["id"], ok(owner.get(f"{API}/auth/me"))["user"]["id"]
    assert commit_import_task.apply(args=[tenant_id, created["id"], user_id]).get() == {"skipped": True}
    assert ok(owner.get(f"{API}/leads"))["total"] == 1
    assert tenant_query(migrator_engine, owner, "SELECT count(*) FROM lead_scores WHERE tenant_id = :t")[0][0] == 1


def test_score_override_is_computed_on_the_server_and_keeps_history(owner: TestClient) -> None:
    import_and_commit(owner, workbook_bytes([lead_row("T-001", "Salon Aurora")]))
    lead = ok(owner.get(f"{API}/leads"))["items"][0]
    url = f"{API}/leads/{lead['id']}/scores"
    full = {"evidence": 30, "relevance": 20, "value": 10, "reachability": 10, "activity": 9}
    assert owner.post(url, json={"components": full}).status_code == 422  # a reason is required
    assert owner.post(url, json={"components": {**full, "evidence": 31}, "reason": "typo"}).status_code == 422
    assert owner.post(url, json={"components": {"evidence": 30}, "reason": "partial"}).status_code == 422
    assert owner.post(url, json={"components": full, "reason": "ok", "total": 100, "tier": "A"}).status_code in (
        201,
        422,
    )
    result = ok(owner.post(url, json={"components": full, "reason": "Re-checked the booking page"}), 201)
    assert (result["total"], result["tier"], result["origin"], result["override_reason"]) == (
        79,
        "B",
        "override",
        "Re-checked the booking page",
    )
    history = ok(owner.get(url))
    assert sum(s["is_current"] for s in history) == 1 and history[0]["is_current"]
    assert {s["rubric_version"] for s in history} == {1}
    rubric = ok(owner.get(f"{API}/rubric"))
    assert [(c["key"], c["max"]) for c in rubric["components"]] == [
        ("evidence", 30),
        ("relevance", 25),
        ("value", 20),
        ("reachability", 15),
        ("activity", 10),
    ]
    assert (rubric["tier_a_min"], rubric["tier_b_min"]) == (80, 60)


# --- the reference workbook ------------------------------------------------------------------

FREE_MAIL = ("abv.bg", "gmail.com", "mail.bg", "yahoo.com", "dir.bg", "hotmail.com", "outlook.com")


@pytest.fixture
def reference_bytes(reference_workbook: Path) -> bytes:
    return reference_workbook.read_bytes()


@pytest.mark.reference_fixture
def test_reference_preview_matches_the_workbook(owner: TestClient, reference_bytes: bytes) -> None:
    created = ok(upload(owner, reference_bytes, "reference.xlsx"), 202)
    report = created["report"]
    assert created["status"] == "ready"
    assert report["sheet"] == "All qualified leads" and len(report["headers"]) == 35
    assert len(report["mapping"]) == 35 and report["unmapped_headers"] == []
    assert report["counts"] == {"create": 94, "rows": 94, "with_errors": 0, "with_warnings": 0}
    assert report["tiers"] == {"A": 21, "B": 62, "C": 11}
    assert report["confidence"] == {"high": 31, "medium": 58, "low": 5}
    assert report["missing_not_found"] == {
        "website_url": 30,
        "phone": 1,
        "email": 58,
        "contact_page": 46,
        "other_channel": 41,
    }
    assert sum(report["missing_not_found"].values()) == 176
    assert report["score_mismatches"] == 0 and report["formula_cells_without_cache"] == 0
    assert report["shortlist"]["rows"] == 25 and report["shortlist"]["unresolved_lead_ids"] == []
    assert report["shortlist"]["creates_leads"] is False and report["summary_sheet"] == "Research summary"
    assert report["issues_sample"] == []
    assert ok(owner.get(f"{API}/leads"))["total"] == 0


@pytest.mark.reference_fixture
def test_reference_import_reproduces_every_documented_figure(
    owner: TestClient, reference_bytes: bytes, migrator_engine: Any
) -> None:
    result = import_and_commit(owner, reference_bytes, "reference.xlsx")["result"]
    assert result["created"] == 94 and "skipped" not in result

    def q(sql: str, **params: Any) -> list[Any]:
        return tenant_query(migrator_engine, owner, sql, **params)

    assert ok(owner.get(f"{API}/leads"))["total"] == 94  # 94 leads, not 119
    assert q("SELECT count(*), count(DISTINCT external_id) FROM leads WHERE tenant_id = :t") == [(94, 94)]
    assert q("SELECT count(*) FROM companies WHERE tenant_id = :t") == [(94,)]

    tiers = dict(q("SELECT tier, count(*) FROM lead_scores WHERE tenant_id = :t AND is_current GROUP BY 1"))
    assert tiers == {"A": 21, "B": 62, "C": 11}
    assert q("SELECT count(*) FROM lead_scores WHERE tenant_id = :t AND total <> source_total") == [(0,)]
    sums = q(
        "SELECT count(*) FROM lead_scores WHERE tenant_id = :t AND total <> (components->>'evidence')::int + "
        "(components->>'relevance')::int + (components->>'value')::int + (components->>'reachability')::int + (components->>'activity')::int"
    )
    assert sums == [(0,)]
    caps = q(
        "SELECT max((components->>'evidence')::int), max((components->>'relevance')::int), max((components->>'value')::int), "
        "max((components->>'reachability')::int), max((components->>'activity')::int) FROM lead_scores WHERE tenant_id = :t"
    )[0]
    assert all(value <= cap for value, cap in zip(caps, (30, 25, 20, 15, 10), strict=True))
    boundary = dict(
        q(
            "SELECT total, min(tier) || max(tier) FROM lead_scores WHERE tenant_id = :t AND total IN (59, 60, 79, 80) GROUP BY 1"
        )
    )
    assert boundary == {59: "CC", 60: "BB", 79: "BB", 80: "AA"}

    confidence = dict(q("SELECT confidence, count(*) FROM lead_assessments WHERE tenant_id = :t GROUP BY 1"))
    assert confidence == {"high": 31, "medium": 58, "low": 5}
    assert q(
        "SELECT DISTINCT checked_on::text, pg_typeof(checked_on)::text FROM lead_assessments WHERE tenant_id = :t"
    ) == [("2026-10-07", "date")]
    assert q(
        "SELECT count(*) FROM lead_assessments WHERE tenant_id = :t AND (source_type, verification_state) <> ('user_import', 'unverified')"
    ) == [(0,)]
    assert q("SELECT count(*) FROM lead_assessments a, jsonb_object_keys(a.raw_row) k WHERE a.tenant_id = :t") == [
        (94 * 35,)
    ]

    listing = dict(q("SELECT listing_id_type, count(*) FROM lead_assessments WHERE tenant_id = :t GROUP BY 1"))
    assert listing == {"place_id": 79, "cid": 15}
    assert q(
        "SELECT count(*) FROM lead_assessments WHERE tenant_id = :t AND listing_id_type = 'cid' AND listing_id !~ '^[0-9]+$'"
    ) == [(0,)]

    assert q("SELECT count(DISTINCT website_status_raw) FROM lead_assessments WHERE tenant_id = :t") == [(14,)]
    assert q("SELECT count(*) FROM lead_assessments WHERE tenant_id = :t AND website_base = 'unknown'") == [(0,)]
    website = dict(q("SELECT website_base, count(*) FROM lead_assessments WHERE tenant_id = :t GROUP BY 1"))
    assert website["not_found"] == 30 and sum(website.values()) == 94

    raw_channels = dict(
        q("SELECT channel_recommendation_raw, count(*) FROM lead_assessments WHERE tenant_id = :t GROUP BY 1")
    )
    assert len(raw_channels) == 4
    preferred = dict(q("SELECT preferred_channel, count(*) FROM lead_assessments WHERE tenant_id = :t GROUP BY 1"))
    assert preferred == {"phone": 66, "email": 28}
    assert q(
        "SELECT preferred_channel, fallback_channels, channel_instruction FROM lead_assessments "
        "WHERE tenant_id = :t AND channel_instruction IS NOT NULL"
    ) == [("phone", [], "ask for the manager")]
    assert q("SELECT count(*) FROM lead_assessments WHERE tenant_id = :t AND fallback_channels = ARRAY['phone']") == [
        (1,)
    ]

    missing = dict(
        q(
            "SELECT k, count(*) FROM lead_assessments a, jsonb_each(a.contact_states) AS s(k, v) "
            "WHERE a.tenant_id = :t AND v->>'state' = 'not_found' AND v->>'raw' = 'Not found' GROUP BY 1"
        )
    )
    assert missing == {"website_url": 30, "phone": 1, "email": 58, "contact_page": 46, "other_channel": 41}
    assert q("SELECT count(*) FROM contact_channels WHERE tenant_id = :t AND lower(raw_value) LIKE '%not found%'") == [
        (0,)
    ]
    assert q("SELECT count(*) FROM companies WHERE tenant_id = :t AND website_url IS NULL") == [(30,)]

    assert q(
        "SELECT count(*) FROM lead_assessments WHERE tenant_id = :t AND raw_row->>'Public business phone' LIKE '%;%'"
    ) == [(17,)]
    purposes = dict(
        q("SELECT purpose, count(*) FROM contact_channels WHERE tenant_id = :t AND kind = 'phone' GROUP BY 1")
    )
    assert purposes["delivery"] == 1 and purposes["emergency"] == 1
    assert q("SELECT allow_sales_use FROM contact_channels WHERE tenant_id = :t AND purpose = 'emergency'") == [
        (False,)
    ]
    assert q("SELECT count(DISTINCT company_id) FROM contact_channels WHERE tenant_id = :t AND kind = 'phone'") == [
        (93,)
    ]
    assert q(
        "SELECT count(*) FROM contact_channels WHERE tenant_id = :t AND kind = 'phone' AND NOT normalized_is_e164"
    ) == [(0,)]

    emails = [
        row[0] for row in q("SELECT normalized_value FROM contact_channels WHERE tenant_id = :t AND kind = 'email'")
    ]
    assert len(emails) == 36
    domains = Counter(email.split("@")[1] for email in emails)
    assert sum(domains[d] for d in FREE_MAIL) == 19
    assert (domains["abv.bg"], domains["gmail.com"], domains["mail.bg"]) == (10, 8, 1)
    assert q(
        "SELECT count(*) FROM leads l JOIN companies c ON c.id = l.company_id WHERE l.tenant_id = :t AND c.website_url IS NULL "
        "AND NOT EXISTS (SELECT 1 FROM contact_channels ch WHERE ch.company_id = c.id AND ch.kind = 'email')"
    ) == [(30,)]

    tags = dict(
        q(
            "SELECT t.name::text, count(*) FROM taggings tg JOIN tags t ON t.id = tg.tag_id WHERE tg.tenant_id = :t GROUP BY 1"
        )
    )
    assert len(tags) == 10 and tags["Other"] == 2
    assert q("SELECT count(*) FROM observations WHERE tenant_id = :t") == [(94,)]
    assert q("SELECT count(*) FROM hypotheses WHERE tenant_id = :t") == [(94,)]
    assert q("SELECT count(DISTINCT outreach_status), min(outreach_status) FROM leads WHERE tenant_id = :t") == [
        (1, "not_contacted")
    ]
    assert (
        q("SELECT count(*) FROM companies WHERE tenant_id = :t AND name ~ '[А-Яа-я]'")[0][0] > 0
    )  # Cyrillic preserved

    # The shortlist is a view over the same leads: the top 25, 21 A plus 4 B, minimum 77.
    assert q("SELECT count(*), min(shortlist_date)::text, min(origin) FROM shortlists WHERE tenant_id = :t") == [
        (1, "2026-10-08", "import")
    ]
    entries = q(
        "SELECT e.rank, l.external_id, e.score_at_snapshot, e.tier_at_snapshot, e.is_filler FROM shortlist_entries e "
        "JOIN leads l ON l.id = e.lead_id WHERE e.tenant_id = :t ORDER BY e.rank"
    )
    assert [e.rank for e in entries] == list(range(1, 26))
    assert Counter(e.tier_at_snapshot for e in entries) == {"A": 21, "B": 4}
    assert sum(e.is_filler for e in entries) == 4 and min(e.score_at_snapshot for e in entries) == 77
    top = q(
        "SELECT l.external_id FROM lead_scores s JOIN leads l ON l.id = s.lead_id WHERE s.tenant_id = :t AND s.is_current "
        "ORDER BY s.total DESC, l.external_id LIMIT 25"
    )
    assert {e.external_id for e in entries} == {row[0] for row in top}
    scores = [e.score_at_snapshot for e in entries]
    assert scores == sorted(scores, reverse=True)  # source rank order is by score


@pytest.mark.reference_fixture
def test_reference_dates_are_stable_in_any_session_time_zone(
    owner: TestClient, reference_bytes: bytes, migrator_engine: Any
) -> None:
    import_and_commit(owner, reference_bytes, "reference.xlsx")
    tenant_id = ok(owner.get(f"{API}/tenant"))["id"]
    for zone in ("UTC", "Europe/Sofia", "America/Los_Angeles", "Pacific/Kiritimati"):
        with migrator_engine.begin() as conn:
            conn.execute(text("SELECT set_config('app.tenant_id', :t, true)"), {"t": tenant_id})
            conn.execute(text(f"SET LOCAL TIME ZONE '{zone}'"))
            values = conn.execute(text("SELECT DISTINCT checked_on::text FROM lead_assessments")).scalars().all()
            sourced = conn.execute(text("SELECT DISTINCT source_date::text FROM contact_channels")).scalars().all()
        assert values == ["2026-10-07"] and sourced == ["2026-10-07"], zone
    sheet = openpyxl.load_workbook(io.BytesIO(owner.get(f"{API}/exports/prospects.xlsx").content))[
        "All qualified leads"
    ]
    column = next(c.column for c in sheet[1] if c.value == "Date checked")
    assert {str(sheet.cell(row=r, column=column).value)[:10] for r in range(2, 96)} == {"2026-10-07"}


@pytest.mark.reference_fixture
def test_reference_reimport_creates_no_duplicates(
    owner: TestClient, reference_bytes: bytes, migrator_engine: Any
) -> None:
    import_and_commit(owner, reference_bytes, "reference.xlsx")
    again = ok(upload(owner, reference_bytes, "reference.xlsx"), 202)
    assert again["report"]["counts"] == {"update": 94, "rows": 94, "with_errors": 0, "with_warnings": 0}
    result = ok(owner.post(f"{API}/imports/{again['id']}/commit"), 202)["result"]
    assert result["unchanged"] == 94 and "created" not in result and "updated" not in result
    for table, expected in (
        ("leads", 94),
        ("companies", 94),
        ("lead_assessments", 94),
        ("lead_scores", 94),
        ("observations", 94),
        ("hypotheses", 94),
        ("shortlists", 1),
        ("shortlist_entries", 25),
    ):
        assert (
            tenant_query(migrator_engine, owner, f"SELECT count(*) FROM {table} WHERE tenant_id = :t")[0][0] == expected
        ), table


@pytest.mark.reference_fixture
def test_reference_export_round_trips_into_another_tenant(
    owner: TestClient, second_tenant: TestClient, reference_bytes: bytes, migrator_engine: Any
) -> None:
    import_and_commit(owner, reference_bytes, "reference.xlsx")
    exported = owner.get(f"{API}/exports/prospects.xlsx").content
    book = openpyxl.load_workbook(io.BytesIO(exported))
    assert book.sheetnames[0] == "All qualified leads" and book.sheetnames[2] == "Research summary"
    assert book.sheetnames[1].startswith("Outreach shortlist 20")  # dated, not "tomorrow"
    leads_sheet = book["All qualified leads"]
    assert [c.value for c in leads_sheet[1]] == HEADERS and leads_sheet.max_row == 95
    assert book[book.sheetnames[1]].max_row == 26
    summary = {row[0]: row[1] for row in book["Research summary"].iter_rows(values_only=True) if row[0]}
    assert (
        summary["Total leads"],
        summary["Tier A (80–100)"],
        summary["Tier B (60–79)"],
        summary["Tier C (below 60)"],
    ) == (94, 21, 62, 11)
    assert (
        summary["Leads with a phone"],
        summary["Leads with an email"],
        summary["Leads with a website recorded"],
    ) == (93, 36, 64)
    assert (summary["Leads on the shortlist"], summary["Shortlist entries below Tier A"]) == (25, 4)

    # Cell-for-cell against the source for the fields that must survive unchanged.
    source = openpyxl.load_workbook(io.BytesIO(reference_bytes), data_only=True)["All qualified leads"]
    position = {header: index for index, header in enumerate(HEADERS)}
    original = {row[0]: row for row in source.iter_rows(min_row=2, values_only=True)}
    round_tripped = {row[0]: row for row in leads_sheet.iter_rows(min_row=2, values_only=True)}
    assert set(original) == set(round_tripped) and len(original) == 94
    exact = [
        h
        for h in HEADERS
        if h not in ("Date checked", "Public business phone", "Other public business contact channel")
    ]
    for lead_id, row in original.items():
        for header in exact:
            assert round_tripped[lead_id][position[header]] == row[position[header]], (lead_id, header)
        assert str(round_tripped[lead_id][position["Date checked"]])[:10] == str(row[position["Date checked"]])[:10]

    result = import_and_commit(second_tenant, exported, "exported.xlsx")["result"]
    assert result["created"] == 94

    def facts(client: TestClient) -> list[Any]:
        return tenant_query(
            migrator_engine,
            client,
            """
            SELECT l.external_id, c.name, c.city, c.website_url, a.checked_on::text, a.confidence, a.website_status_raw,
                   a.listing_id_type, a.listing_id, a.preferred_channel, a.channel_instruction, a.outreach_opening,
                   a.discovery_question, a.notes, a.contact_states, s.components, s.total, s.tier,
                   (SELECT array_agg(ch.kind || ':' || ch.normalized_value || ':' || ch.purpose ORDER BY ch.kind, ch.normalized_value)
                    FROM contact_channels ch WHERE ch.company_id = c.id),
                   (SELECT o.text FROM observations o WHERE o.lead_id = l.id), (SELECT h.text FROM hypotheses h WHERE h.lead_id = l.id),
                   (SELECT array_agg(t.name::text ORDER BY t.name::text) FROM taggings tg JOIN tags t ON t.id = tg.tag_id WHERE tg.entity_id = l.id)
            FROM leads l JOIN companies c ON c.id = l.company_id
            JOIN lead_assessments a ON a.lead_id = l.id AND a.is_current
            JOIN lead_scores s ON s.lead_id = l.id AND s.is_current
            WHERE l.tenant_id = :t ORDER BY l.external_id
            """,
        )

    first, second = facts(owner), facts(second_tenant)
    assert len(first) == 94 and first == second
